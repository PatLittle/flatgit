#!/usr/bin/env python3
"""Generate FlatGit demonstration files in test-content/generated."""

from __future__ import annotations

import asyncio
import csv
import json
import os
import shutil
import sqlite3
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import duckdb
import markdown
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from playwright.async_api import async_playwright
from weasyprint import HTML


SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent
OUTPUT = SCRIPT_DIR / "generated"
PAGE_URL = os.environ.get("FLATGIT_PAGE_URL", "https://patlittle.github.io/flatgit/")

ROWS = [
    {"id": 1, "name": "Markdown guide", "category": "documentation", "count": 14, "price": 0, "date": "2026-01-15", "active": True, "url": "https://github.com/PatLittle/flatgit"},
    {"id": 2, "name": "Orders table", "category": "data", "count": 240, "price": 19.95, "date": "2026-02-03", "active": True, "url": "https://patlittle.github.io/flatgit/"},
    {"id": 3, "name": "Office preview", "category": "document", "count": 8, "price": 42.5, "date": "2026-03-22", "active": False, "url": "https://view.officeapps.live.com/"},
    {"id": 4, "name": "SQLite query", "category": "database", "count": 86, "price": 7.25, "date": "2026-04-10", "active": True, "url": "https://sqlite.org/"},
    {"id": 5, "name": "Mermaid diagram", "category": "diagram", "count": 31, "price": 12, "date": "2026-05-19", "active": True, "url": "https://mermaid.js.org/"},
    {"id": 6, "name": "Parquet sample", "category": "data", "count": 175, "price": 33.4, "date": "2026-06-07", "active": False, "url": "https://parquet.apache.org/"},
]


def run(command: list[str], cwd: Path = ROOT) -> None:
    subprocess.run(command, cwd=cwd, check=True)


def write_text_examples() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    fields = list(ROWS[0])
    for name, delimiter in (("demo.csv", ","), ("demo.tsv", "\t")):
        with (OUTPUT / name).open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields, delimiter=delimiter)
            writer.writeheader()
            writer.writerows(ROWS)

    nested = [
        {**row, "details": {"label": row["name"].lower().replace(" ", "-"), "score": row["count"] / 10}}
        for row in ROWS
    ]
    (OUTPUT / "demo.json").write_text(json.dumps(nested, indent=2) + "\n", encoding="utf-8")
    (OUTPUT / "demo.jsonl").write_text("\n".join(json.dumps(row) for row in nested) + "\n", encoding="utf-8")

    geojson = {
        "type": "FeatureCollection",
        "features": [
            {"type": "Feature", "properties": {"name": "Ottawa", "status": "active"}, "geometry": {"type": "Point", "coordinates": [-75.6972, 45.4215]}},
            {"type": "Feature", "properties": {"name": "Toronto", "status": "planned"}, "geometry": {"type": "Point", "coordinates": [-79.3832, 43.6532]}},
            {"type": "Feature", "properties": {"name": "Montreal", "status": "complete"}, "geometry": {"type": "Point", "coordinates": [-73.5673, 45.5019]}},
        ],
    }
    (OUTPUT / "demo.geojson").write_text(json.dumps(geojson, indent=2) + "\n", encoding="utf-8")

    (OUTPUT / "demo.md").write_text(
        """# FlatGit Markdown demonstration

This file demonstrates **GitHub-Flavoured Markdown**, relative links, tables, code, and Mermaid.

| Viewer | Example | Interactive |
| --- | --- | --- |
| Table | CSV and TSV | Yes |
| Document | Markdown | Yes |
| Database | SQLite | Yes |

```python
def greet(name: str) -> str:
    return f\"Hello, {name}\"
```

```mermaid
flowchart LR
    A[GitHub file] --> B[FlatGit]
    B --> C[Interactive preview]
```

[Open the generated CSV](demo.csv)
""",
        encoding="utf-8",
    )
    (OUTPUT / "demo.mmd").write_text(
        """flowchart TD
    A[Paste a public GitHub URL] --> B{File type}
    B -->|CSV JSON Parquet| C[Interactive table]
    B -->|Markdown Mermaid| D[Rendered document]
    B -->|SQLite| E[Read-only query]
    B -->|Code or text| F[Syntax viewer]
""",
        encoding="utf-8",
    )
    (OUTPUT / "demo.txt").write_text(
        "FlatGit plain-text fallback\n\nUnknown text extensions remain readable even when no dedicated syntax mode exists.\n",
        encoding="utf-8",
    )
    (OUTPUT / "demo.py").write_text(
        """from collections import Counter


def categories(rows):
    \"\"\"Count records by category.\"\"\"
    return Counter(row[\"category\"] for row in rows)


if __name__ == \"__main__\":
    print(categories([{"category": "data"}, {"category": "data"}, {"category": "docs"}]))
""",
        encoding="utf-8",
    )
    (OUTPUT / "demo.html").write_text(
        """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>FlatGit HTML source</title></head>
<body><main><h1>HTML shown as source</h1><p>FlatGit safely displays this file in its code viewer.</p></main></body></html>
""",
        encoding="utf-8",
    )
    (OUTPUT / "demo.svg").write_text(
        """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 800 420" role="img" aria-labelledby="title desc">
  <title id="title">FlatGit sample values</title><desc id="desc">A simple bar chart with five values.</desc>
  <rect width="800" height="420" fill="#f6f8fa"/>
  <text x="50" y="55" font-family="Arial, sans-serif" font-size="28" font-weight="700" fill="#1f2328">FlatGit SVG preview</text>
  <g fill="#0969da">
    <rect x="70" y="250" width="90" height="110"/><rect x="205" y="190" width="90" height="170"/>
    <rect x="340" y="120" width="90" height="240"/><rect x="475" y="220" width="90" height="140"/>
    <rect x="610" y="155" width="90" height="205"/>
  </g>
  <g font-family="Arial, sans-serif" font-size="18" fill="#656d76" text-anchor="middle">
    <text x="115" y="392">CSV</text><text x="250" y="392">JSON</text><text x="385" y="392">Markdown</text><text x="520" y="392">SQLite</text><text x="655" y="392">Parquet</text>
  </g>
</svg>
""",
        encoding="utf-8",
    )


def write_database_and_parquet() -> None:
    database = OUTPUT / "demo.sqlite"
    database.unlink(missing_ok=True)
    connection = sqlite3.connect(database)
    connection.execute("CREATE TABLE items (id INTEGER PRIMARY KEY, name TEXT, category TEXT, count INTEGER, price REAL, date TEXT, active INTEGER, url TEXT)")
    connection.executemany(
        "INSERT INTO items VALUES (:id, :name, :category, :count, :price, :date, :active, :url)",
        [{**row, "active": int(row["active"])} for row in ROWS],
    )
    connection.execute("CREATE VIEW category_summary AS SELECT category, COUNT(*) AS records, SUM(count) AS total_count FROM items GROUP BY category")
    connection.commit()
    connection.close()

    parquet = OUTPUT / "demo.parquet"
    parquet.unlink(missing_ok=True)
    db = duckdb.connect()
    db.execute("CREATE TABLE items (id INTEGER, name VARCHAR, category VARCHAR, count INTEGER, price DOUBLE, date DATE, active BOOLEAN, url VARCHAR)")
    db.executemany("INSERT INTO items VALUES (?, ?, ?, ?, ?, ?, ?, ?)", [tuple(row.values()) for row in ROWS])
    db.execute(f"COPY items TO '{parquet.as_posix()}' (FORMAT PARQUET, COMPRESSION ZSTD)")
    db.close()


def write_excel() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FlatGit demo"
    headers = list(ROWS[0])
    sheet.append(headers)
    for row in ROWS:
        sheet.append([row[column] for column in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0969DA")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {"A": 8, "B": 24, "C": 18, "D": 12, "E": 12, "F": 14, "G": 10, "H": 44}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    workbook.save(OUTPUT / "demo.xlsx")


def write_readme_docx() -> None:
    if not shutil.which("pandoc"):
        raise RuntimeError("Pandoc is required to generate README.docx")
    run([
        "pandoc", str(ROOT / "README.md"), "--from=gfm", "--to=docx", "--toc",
        "--metadata", "title=FlatGit documentation", "--output", str(OUTPUT / "README.docx"),
    ])


def write_readme_pdf() -> None:
    source = (ROOT / "README.md").read_text(encoding="utf-8")
    body = markdown.markdown(source, extensions=["fenced_code", "tables", "toc"])
    stylesheet = """
      @page { size: Letter; margin: 0.65in 0.7in 0.7in; @bottom-right { content: "Page " counter(page) " of " counter(pages); font: 8pt Arial; color: #656d76; } }
      body { font-family: Arial, sans-serif; color: #1f2328; font-size: 9.5pt; line-height: 1.45; }
      h1 { font-size: 25pt; margin: 0 0 8pt; color: #0d1117; }
      h2 { font-size: 16pt; margin: 18pt 0 6pt; color: #0d1117; break-after: avoid; }
      h3 { font-size: 12pt; margin: 13pt 0 5pt; color: #0d1117; break-after: avoid; }
      p, li { orphans: 3; widows: 3; }
      a { color: #0969da; text-decoration: none; }
      code { font-family: "DejaVu Sans Mono", monospace; font-size: 8.3pt; background: #f6f8fa; padding: 1pt 2pt; overflow-wrap: anywhere; }
      pre { padding: 8pt; background: #f6f8fa; border: 0.5pt solid #d0d7de; border-radius: 4pt; white-space: pre-wrap; overflow-wrap: anywhere; break-inside: avoid; }
      table { width: 100%; border-collapse: collapse; table-layout: fixed; margin: 7pt 0 10pt; font-size: 7.5pt; }
      th, td { border: 0.5pt solid #d0d7de; padding: 4pt; vertical-align: top; overflow-wrap: anywhere; }
      th { color: white; background: #24292f; text-align: left; }
      tr:nth-child(even) td { background: #f6f8fa; }
      blockquote { margin-left: 0; padding-left: 10pt; border-left: 3pt solid #d0d7de; color: #656d76; }
    """
    document = f"<!doctype html><html><head><meta charset='utf-8'><style>{stylesheet}</style></head><body>{body}</body></html>"
    HTML(string=document, base_url=str(ROOT)).write_pdf(OUTPUT / "README.pdf")


async def capture_screenshots() -> None:
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch()
        for theme, filename, image_type in (
            ("light", "flatgit-light.jpg", "jpeg"),
            ("dark", "flatgit-dark.png", "png"),
        ):
            context = await browser.new_context(viewport={"width": 1440, "height": 900}, device_scale_factor=1)
            page = await context.new_page()
            await page.add_init_script(f"localStorage.setItem('flatgit-theme', '{theme}')")
            await page.goto(PAGE_URL, wait_until="networkidle")
            await page.wait_for_selector("#welcome:not(.hidden)")
            await page.screenshot(path=OUTPUT / filename, type=image_type, quality=90 if image_type == "jpeg" else None, full_page=False)
            await context.close()
        await browser.close()


def presentation_data() -> dict:
    soup = BeautifulSoup((ROOT / "index.html").read_text(encoding="utf-8"), "html.parser")
    features = [
        {"title": card.find("strong").get_text(" ", strip=True), "description": card.find("span").get_text(" ", strip=True)}
        for card in soup.select(".feature-grid article")
    ]
    formats = [row.find("td").get_text(" ", strip=True) for row in soup.select(".usage-table tbody tr")]
    return {
        "title": soup.select_one(".hero-card h1").get_text(" ", strip=True),
        "lede": soup.select_one(".hero-card .lede").get_text(" ", strip=True),
        "features": features,
        "formats": formats,
        "lightScreenshot": str((OUTPUT / "flatgit-light.jpg").resolve()),
        "darkScreenshot": str((OUTPUT / "flatgit-dark.png").resolve()),
        "output": str((OUTPUT / "flatgit-demo.pptx").resolve()),
    }


def write_presentation() -> None:
    data_path = OUTPUT / "presentation-data.json"
    data_path.write_text(json.dumps(presentation_data(), indent=2) + "\n", encoding="utf-8")
    run(["node", str(SCRIPT_DIR / "build_presentation.mjs"), str(data_path)], cwd=SCRIPT_DIR)
    data_path.unlink()
    normalize_pptx_package(OUTPUT / "flatgit-demo.pptx")


def normalize_pptx_package(path: Path) -> None:
    """Remove PptxGenJS content-type entries that do not have package parts."""
    ElementTree.register_namespace("", "http://schemas.openxmlformats.org/package/2006/content-types")
    with zipfile.ZipFile(path) as source:
        members = {info.filename for info in source.infolist()}
        content_types = ElementTree.fromstring(source.read("[Content_Types].xml"))
        namespace = "{http://schemas.openxmlformats.org/package/2006/content-types}"
        for override in list(content_types.findall(f"{namespace}Override")):
            target = override.attrib.get("PartName", "").lstrip("/")
            if target and target not in members:
                content_types.remove(override)

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as stream:
            temporary = Path(stream.name)
        try:
            with zipfile.ZipFile(temporary, "w") as destination:
                for info in source.infolist():
                    payload = source.read(info.filename)
                    if info.filename == "[Content_Types].xml":
                        payload = ElementTree.tostring(content_types, encoding="utf-8", xml_declaration=True)
                    destination.writestr(info, payload)
            temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)


def write_generated_index() -> None:
    files = [
        ("README.docx", "Word conversion of the project README"),
        ("README.pdf", "PDF conversion of the project README"),
        ("flatgit-demo.pptx", "Presentation based on the FlatGit landing page"),
        ("flatgit-light.jpg", "Landing-page screenshot in light mode"),
        ("flatgit-dark.png", "Landing-page screenshot in dark mode"),
        ("demo.csv", "CSV interactive-table sample"),
        ("demo.tsv", "TSV interactive-table sample"),
        ("demo.json", "Nested JSON sample"),
        ("demo.jsonl", "JSON Lines sample"),
        ("demo.geojson", "GeoJSON sample"),
        ("demo.md", "Markdown with a table, code, relative link, and Mermaid"),
        ("demo.mmd", "Standalone Mermaid diagram"),
        ("demo.sqlite", "SQLite database with a table and summary view"),
        ("demo.parquet", "Parquet table sample"),
        ("demo.xlsx", "Excel workbook sample"),
        ("demo.svg", "Sanitized SVG preview sample"),
        ("demo.py", "Syntax highlighting and folding sample"),
        ("demo.html", "HTML source-view sample"),
        ("demo.txt", "Plain-text fallback sample"),
    ]
    lines = ["# Generated FlatGit test content", "", "These files are regenerated by `test-content/generate_test_content.py`.", ""]
    for filename, description in files:
        flatgit = f"https://patlittle.github.io/flatgit/full/PatLittle/flatgit/blob/main/test-content/generated/{filename}"
        lines.append(f"- [`{filename}`]({filename}) - {description}. [Open in FlatGit]({flatgit})")
    (OUTPUT / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    if OUTPUT.exists():
        shutil.rmtree(OUTPUT)
    OUTPUT.mkdir(parents=True)
    write_text_examples()
    write_database_and_parquet()
    write_excel()
    write_readme_docx()
    write_readme_pdf()
    asyncio.run(capture_screenshots())
    write_presentation()
    write_generated_index()
    print(f"Generated {len(list(OUTPUT.iterdir()))} files in {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
